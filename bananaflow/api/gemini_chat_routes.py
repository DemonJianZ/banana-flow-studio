"""
gemini_chat_routes.py — Agent 对话面板后端路由

• POST /api/gemini-chat/upload   — 上传附件（图片/视频/文档），返回 LLM 可用内容
• POST /api/gemini-chat/chat     — 流式对话（支持多模态附件）
• DELETE /api/gemini-chat/chat/{thread_id} — 清空会话历史
"""
from __future__ import annotations

import base64
import io
import json
import os
import threading
import uuid
from collections import defaultdict
from pathlib import Path
from typing import Any, Dict, List, Optional

import httpx
from fastapi import APIRouter, File, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

try:
    from ..core.config import MODEL_AGENT_CHAT
    from ..core.logging import sys_logger
except Exception:
    from core.config import MODEL_AGENT_CHAT
    from core.logging import sys_logger

gemini_chat_router = APIRouter(prefix="/api/gemini-chat", tags=["gemini-chat"])

# ── 配置 ───────────────────────────────────────────────────────────────────────

_SYSTEM_PROMPT = (
    "你是 AI 小禹智能体，用户的 AI 创意伙伴，专注于电商图片生成、视频创作和故事分镜设计。\n"
    "请用中文回复，保持专业、简洁、富有创意。\n"
    "当用户上传图片时，仔细分析图片内容并给出专业建议。\n"
    "当用户上传文档时，认真阅读文档内容并根据内容回答问题。\n"
    "当用户询问能做什么时，介绍你可以帮助生成电商主图、短视频分镜、产品三视图、场景广告等。"
)

_MAX_HISTORY_TURNS = 20
_MAX_IMAGE_BYTES = 8 * 1024 * 1024   # 8 MB
_MAX_DOC_TEXT_CHARS = 12_000          # 截断过长文档
_IMG_MAX_PX = 1536                    # 图片长边最大像素

# ── 会话历史 ───────────────────────────────────────────────────────────────────

_history_store: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
_history_lock = threading.Lock()


def _get_history(thread_id: str) -> List[Dict[str, Any]]:
    with _history_lock:
        return list(_history_store[thread_id])


def _append_history(thread_id: str, role: str, content: Any) -> None:
    with _history_lock:
        _history_store[thread_id].append({"role": role, "content": content})
        if len(_history_store[thread_id]) > _MAX_HISTORY_TURNS:
            _history_store[thread_id] = _history_store[thread_id][-_MAX_HISTORY_TURNS:]


def _clear_history(thread_id: str) -> None:
    with _history_lock:
        _history_store.pop(thread_id, None)


# ── 文件解析工具 ───────────────────────────────────────────────────────────────

def _resize_image_bytes(data: bytes, mime: str) -> tuple[bytes, str]:
    """用 Pillow 压缩图片到 _IMG_MAX_PX 长边，返回 (bytes, mime)。"""
    try:
        from PIL import Image
        img = Image.open(io.BytesIO(data))
        if max(img.size) > _IMG_MAX_PX:
            ratio = _IMG_MAX_PX / max(img.size)
            new_size = (int(img.width * ratio), int(img.height * ratio))
            img = img.resize(new_size, Image.LANCZOS)
        # 统一输出 JPEG（除 PNG 透明图外）
        fmt = "PNG" if img.mode in ("RGBA", "LA", "P") else "JPEG"
        buf = io.BytesIO()
        img.convert("RGB" if fmt == "JPEG" else img.mode).save(buf, format=fmt, quality=85)
        return buf.getvalue(), f"image/{fmt.lower()}"
    except Exception:
        return data, mime


def _extract_pdf_text(data: bytes) -> str:
    try:
        import PyPDF2
        reader = PyPDF2.PdfReader(io.BytesIO(data))
        parts = []
        for page in reader.pages:
            text = page.extract_text() or ""
            if text.strip():
                parts.append(text.strip())
        return "\n\n".join(parts)
    except Exception as e:
        return f"[PDF 解析失败: {e}]"


def _extract_docx_text(data: bytes) -> str:
    try:
        import docx
        doc = docx.Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    except Exception as e:
        return f"[Word 解析失败: {e}]"


def _extract_xlsx_text(data: bytes) -> str:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts = []
        for sheet in wb.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                cells = [str(c) if c is not None else "" for c in row]
                if any(c.strip() for c in cells):
                    rows.append("\t".join(cells))
            if rows:
                parts.append(f"[Sheet: {sheet.title}]\n" + "\n".join(rows))
        return "\n\n".join(parts)
    except Exception as e:
        return f"[Excel 解析失败: {e}]"


def _extract_pptx_text(data: bytes) -> str:
    try:
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        slides = []
        for i, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        t = para.text.strip()
                        if t:
                            texts.append(t)
            if texts:
                slides.append(f"[第 {i} 页]\n" + "\n".join(texts))
        return "\n\n".join(slides)
    except Exception as e:
        return f"[PPT 解析失败: {e}]"


_DOC_EXTRACTORS = {
    ".pdf":  _extract_pdf_text,
    ".docx": _extract_docx_text,
    ".doc":  _extract_docx_text,
    ".xlsx": _extract_xlsx_text,
    ".xls":  _extract_xlsx_text,
    ".pptx": _extract_pptx_text,
    ".ppt":  _extract_pptx_text,
}

_IMAGE_EXTS  = {".jpg", ".jpeg", ".png", ".gif", ".webp", ".bmp", ".tiff"}
_VIDEO_EXTS  = {".mp4", ".mov", ".avi", ".mkv", ".webm", ".flv", ".wmv"}


# ── 上传接口 ───────────────────────────────────────────────────────────────────

@gemini_chat_router.post("/upload")
async def upload_attachment(file: UploadFile = File(...)):
    """
    解析上传的附件，返回 LLM 可消费的内容。
    返回格式：
      {
        "name": "file.pdf",
        "file_type": "image" | "video" | "doc" | "unknown",
        "mime": "...",
        "data_url": "data:image/jpeg;base64,...",  # 仅 image
        "text_content": "...",                     # 仅 doc
        "error": "..."                             # 解析失败时
      }
    """
    name = str(file.filename or "file").strip()
    ext  = Path(name).suffix.lower()
    mime = str(file.content_type or "application/octet-stream")

    data = await file.read()
    if not data:
        return {"name": name, "file_type": "unknown", "mime": mime, "error": "文件内容为空"}

    # ── 图片 ──
    if ext in _IMAGE_EXTS or mime.startswith("image/"):
        if len(data) > _MAX_IMAGE_BYTES:
            return {"name": name, "file_type": "image", "mime": mime,
                    "error": f"图片超过 {_MAX_IMAGE_BYTES // (1024*1024)} MB 限制"}
        data, mime = _resize_image_bytes(data, mime)
        b64 = base64.b64encode(data).decode()
        data_url = f"data:{mime};base64,{b64}"
        return {"name": name, "file_type": "image", "mime": mime, "data_url": data_url}

    # ── 视频 ──
    if ext in _VIDEO_EXTS or mime.startswith("video/"):
        return {"name": name, "file_type": "video", "mime": mime}

    # ── 文档 ──
    if ext in _DOC_EXTRACTORS:
        text = _DOC_EXTRACTORS[ext](data)
        if len(text) > _MAX_DOC_TEXT_CHARS:
            text = text[:_MAX_DOC_TEXT_CHARS] + f"\n…（内容已截断，原文 {len(text)} 字符）"
        return {"name": name, "file_type": "doc", "mime": mime, "text_content": text}

    return {"name": name, "file_type": "unknown", "mime": mime}


# ── 对话接口 ───────────────────────────────────────────────────────────────────

class AttachmentIn(BaseModel):
    file_type: str            # "image" | "video" | "doc" | "unknown"
    name: str
    mime: Optional[str] = ""
    data_url: Optional[str] = None   # image: data:image/...;base64,...
    text_content: Optional[str] = None  # doc text


class ChatRequest(BaseModel):
    message: str
    thread_id: str = "default"
    attachments: List[AttachmentIn] = []


def _get_api_key() -> str:
    return str(os.getenv("DEEPSEEK_API_KEY", "")).strip()

def _get_base_url() -> str:
    return str(os.getenv("DEEPSEEK_BASE_URL", "https://api.deepseek.com")).strip().rstrip("/")


def _build_user_content(message: str, attachments: List[AttachmentIn]) -> Any:
    """
    将用户消息 + 附件组合成 LLM content 字段。
    - 若有图片：返回 list（OpenAI vision 格式）
    - 若有文档：将文档文本拼入纯文本消息
    - 若仅文本：返回 str
    """
    doc_blocks: List[str] = []
    image_parts: List[Dict[str, Any]] = []
    video_notes: List[str] = []

    for att in attachments:
        if att.file_type == "image" and att.data_url:
            image_parts.append({
                "type": "image_url",
                "image_url": {"url": att.data_url, "detail": "high"},
            })
        elif att.file_type == "doc" and att.text_content:
            doc_blocks.append(
                f"【附件文件：{att.name}】\n{att.text_content}"
            )
        elif att.file_type == "video":
            video_notes.append(f"（用户上传了视频文件：{att.name}）")
        elif att.file_type == "unknown":
            doc_blocks.append(f"（用户上传了文件：{att.name}）")

    # 拼接文本部分
    text_parts: List[str] = []
    if doc_blocks:
        text_parts.append("\n\n---\n".join(doc_blocks))
    if video_notes:
        text_parts.append(" ".join(video_notes))
    if message.strip():
        text_parts.append(message.strip())
    full_text = "\n\n".join(text_parts) if text_parts else message

    if image_parts:
        content: List[Dict[str, Any]] = [{"type": "text", "text": full_text}]
        content.extend(image_parts)
        return content

    return full_text


@gemini_chat_router.post("/chat")
async def chat(req: ChatRequest):
    """
    流式对话接口。SSE 格式：
      data: {"type": "token",  "content": "..."}
      data: {"type": "done"}
      data: {"type": "error", "content": "..."}
    """
    thread_id    = str(req.thread_id or "default").strip() or "default"
    user_message = str(req.message or "").strip()
    has_attachments = bool(req.attachments)

    if not user_message and not has_attachments:
        async def _empty():
            yield f"data: {json.dumps({'type': 'error', 'content': '消息不能为空'})}\n\n"
        return StreamingResponse(_empty(), media_type="text/event-stream")

    history = _get_history(thread_id)
    user_content = _build_user_content(user_message, req.attachments)

    messages: List[Dict[str, Any]] = [{"role": "system", "content": _SYSTEM_PROMPT}]
    messages.extend(history)
    messages.append({"role": "user", "content": user_content})

    api_key  = _get_api_key()
    base_url = _get_base_url()
    raw_model = str(MODEL_AGENT_CHAT or "deepseek-v4-flash").strip()
    model_name = raw_model.split(":", 1)[-1].strip() if ":" in raw_model else raw_model

    async def _stream():
        collected: List[str] = []
        yield f"data: {json.dumps({'type': 'ping'})}\n\n"

        if not api_key:
            yield f"data: {json.dumps({'type': 'error', 'content': 'DEEPSEEK_API_KEY 未配置'})}\n\n"
            return

        url = f"{base_url}/chat/completions"
        headers = {
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json",
            "Accept": "text/event-stream",
        }
        payload: Dict[str, Any] = {
            "model": model_name,
            "messages": messages,
            "stream": True,
        }

        try:
            async with httpx.AsyncClient(timeout=120.0, trust_env=False) as client:
                async with client.stream("POST", url, headers=headers, json=payload) as resp:
                    if resp.status_code != 200:
                        body = await resp.aread()
                        err = body.decode(errors="replace")[:400]
                        sys_logger.warning(f"[gemini_chat] DeepSeek HTTP {resp.status_code}: {err}")
                        yield f"data: {json.dumps({'type': 'error', 'content': f'API 错误 {resp.status_code}: {err}'})}\n\n"
                        return

                    async for line in resp.aiter_lines():
                        line = line.strip()
                        if not line or line == "data: [DONE]":
                            continue
                        if line.startswith("data:"):
                            try:
                                chunk = json.loads(line[5:].strip())
                            except Exception:
                                continue
                            delta = (chunk.get("choices", [{}])[0].get("delta", {}).get("content") or "")
                            if delta:
                                collected.append(delta)
                                yield f"data: {json.dumps({'type': 'token', 'content': delta})}\n\n"

        except httpx.TimeoutException:
            sys_logger.warning("[gemini_chat] DeepSeek request timed out")
            yield f"data: {json.dumps({'type': 'error', 'content': '请求超时，请稍后重试'})}\n\n"
            return
        except Exception as exc:
            sys_logger.exception(f"[gemini_chat] streaming error: {exc}")
            yield f"data: {json.dumps({'type': 'error', 'content': str(exc)})}\n\n"
            return

        reply = "".join(collected)
        if reply:
            # 历史中只保存文本版本（避免 base64 撑爆内存）
            history_text = user_message or "[附件消息]"
            _append_history(thread_id, "user", history_text)
            _append_history(thread_id, "assistant", reply)
            sys_logger.info(
                f"[gemini_chat] thread={thread_id} model={model_name} "
                f"attachments={len(req.attachments)} reply_len={len(reply)}"
            )

        yield f"data: {json.dumps({'type': 'done'})}\n\n"

    return StreamingResponse(
        _stream(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@gemini_chat_router.delete("/chat/{thread_id}")
async def clear_chat(thread_id: str):
    _clear_history(thread_id)
    return {"ok": True, "thread_id": thread_id}
